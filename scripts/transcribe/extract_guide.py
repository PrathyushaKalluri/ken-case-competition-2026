"""Extract plain text from the Ken 2026 household interview guide docx."""
import docx
import sys
from pathlib import Path

SRC = Path("output/docx/Ken_2026_Household_Interview_Guide.docx")
OUT = Path("scripts/transcribe/interview_guide.txt")

def extract(path: Path) -> str:
    d = docx.Document(str(path))
    lines = []
    for para in d.paragraphs:
        t = para.text.strip()
        if t:
            lines.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)

if __name__ == "__main__":
    text = extract(SRC)
    OUT.write_text(text, encoding="utf-8")
    print(f"Wrote {len(text)} chars to {OUT}")
    print("---- preview ----")
    print(text[:2000])
